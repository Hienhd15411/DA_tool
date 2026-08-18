#!/usr/bin/env python3
"""Generate file bài tập riêng cho học viên (xlsx + csv).

Cấu trúc 5 cột theo yêu cầu:
  Tên bài | Độ khó | Thời gian | Skill | Business context

Reuse parsing từ gen_student_curriculum.py.
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def clean_md(s: str) -> str:
    """Strip markdown formatting trong title/skill (**bold**, `code`)."""
    if not s:
        return ""
    s = re.sub(r"\*\*(.*?)\*\*", r"\1", s)
    s = re.sub(r"`(.*?)`", r"\1", s)
    s = re.sub(r"\*(.*?)\*", r"\1", s)
    return s.strip()

# Reuse parser
sys.path.insert(0, str(Path(__file__).parent))
from gen_student_curriculum import parse_all  # noqa: E402

ROOT = Path(__file__).parent.parent
OUT_XLSX = ROOT / "docs" / "bai_tap_hoc_vien.xlsx"
OUT_CSV  = ROOT / "docs" / "bai_tap_hoc_vien.csv"


HEADER_FILL = PatternFill("solid", fgColor="F97316")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(border_style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_workbook(exercises):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bài tập"

    headers = ["Tên bài", "Độ khó", "Thời gian", "Skill", "Business context"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER

    for ex in exercises:
        ws.append([
            f"{ex['id']} — {clean_md(ex['title'])}",
            f"L{ex['level']}",
            ex["est"],
            clean_md(ex["skill"]),
            ex["business_context"],
        ])

    n = len(exercises)
    for row in range(2, n + 2):
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=row, column=col)
            c.alignment = WRAP
            c.border = BORDER
        ws.row_dimensions[row].height = 90

    widths = [62, 8, 12, 36, 75]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    return wb


def write_csv(path, exercises):
    headers = ["Tên bài", "Độ khó", "Thời gian", "Skill", "Business context"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for ex in exercises:
            w.writerow([
                f"{ex['id']} — {clean_md(ex['title'])}",
                f"L{ex['level']}",
                ex["est"],
                clean_md(ex["skill"]),
                ex["business_context"],
            ])


def main():
    exercises = parse_all()
    print(f"Parsed {len(exercises)} exercises")

    wb = build_workbook(exercises)
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    write_csv(OUT_CSV, exercises)

    print(f"✅ Done")
    print(f"   {OUT_XLSX} ({OUT_XLSX.stat().st_size // 1024} KB)")
    print(f"   {OUT_CSV} ({OUT_CSV.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
